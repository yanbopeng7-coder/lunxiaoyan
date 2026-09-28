# -*- coding: utf-8 -*-
"""论小研 · AI 科研论文辅助平台"""
from __future__ import print_function

import json
import os
import time
import uuid

from flask import Flask, Response, jsonify, render_template, request, send_from_directory
from werkzeug.utils import secure_filename

try:
    from dotenv import load_dotenv

    load_dotenv()
except Exception:
    pass

from engine import DEMO_THESIS, health_status, parse_bytes, run_module


def ensure_demo_docx():
    path = os.path.join(DEMO_DIR, u"计算机专业本科毕业论文初稿.docx")
    if os.path.isfile(path):
        return
    try:
        from docx import Document

        doc = Document()
        doc.add_heading(u"计算机专业本科毕业论文初稿", level=0)
        for line in DEMO_THESIS.splitlines():
            doc.add_paragraph(line)
        doc.save(path)
    except Exception:
        pass

ROOT = os.path.dirname(os.path.abspath(__file__))
UPLOAD_DIR = os.path.join(ROOT, "data", "uploads")
HISTORY_DIR = os.path.join(ROOT, "data", "history")
DEMO_DIR = os.path.join(ROOT, "static", "demo")

for d in (UPLOAD_DIR, HISTORY_DIR, DEMO_DIR):
    if not os.path.isdir(d):
        os.makedirs(d)

ensure_demo_docx()

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
app.config["TEMPLATES_AUTO_RELOAD"] = True
app.jinja_env.auto_reload = True
app.config["SEND_FILE_MAX_AGE_DEFAULT"] = 0


@app.after_request
def _no_cache(resp):
    resp.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    resp.headers["Pragma"] = "no-cache"
    return resp


def _json(data, code=200):
    return Response(
        json.dumps(data, ensure_ascii=False),
        status=code,
        mimetype="application/json; charset=utf-8",
    )


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/api/health")
def api_health():
    return _json(health_status())


@app.route("/api/demo-thesis")
def api_demo_thesis():
    return _json(
        {
            "filename": u"计算机专业本科毕业论文初稿.docx",
            "text": DEMO_THESIS,
        }
    )


@app.route("/static/demo/<path:name>")
def demo_file(name):
    return send_from_directory(DEMO_DIR, name, as_attachment=True)


@app.route("/api/parse", methods=["POST"])
def api_parse():
    f = request.files.get("file")
    if not f or not f.filename:
        return _json({"ok": False, "error": u"请选择文件"}, 400)
    filename = f.filename
    data = f.read()
    try:
        text = parse_bytes(filename, data)
    except Exception as e:
        return _json({"ok": False, "error": str(e)}, 400)
    safe = secure_filename(filename) or "upload.txt"
    saved = "%s_%s" % (int(time.time()), safe)
    path = os.path.join(UPLOAD_DIR, saved)
    with open(path, "wb") as out:
        out.write(data)
    return _json({"ok": True, "filename": filename, "text": text, "chars": len(text)})


@app.route("/api/run", methods=["POST"])
def api_run():
    body = request.get_json(silent=True) or {}
    action = body.get("action") or "chat"
    try:
        result = run_module(action, body)
    except Exception as e:
        return _json({"ok": False, "error": str(e)}, 500)
    result["ok"] = True
    result["action"] = action
    return _json(result)


@app.route("/api/history", methods=["GET", "POST", "DELETE"])
def api_history():
    path = os.path.join(HISTORY_DIR, "history.json")
    records = []
    if os.path.isfile(path):
        try:
            with open(path, "r", encoding="utf-8") as f:
                records = json.load(f)
        except Exception:
            records = []
    if request.method == "GET":
        module = request.args.get("module")
        if module:
            records = [r for r in records if r.get("module") == module]
        return _json({"ok": True, "records": records[:80]})
    if request.method == "DELETE":
        hid = (request.get_json(silent=True) or {}).get("id")
        if hid:
            records = [r for r in records if r.get("id") != hid]
        else:
            records = []
        with open(path, "w", encoding="utf-8") as f:
            json.dump(records, f, ensure_ascii=False, indent=2)
        return _json({"ok": True})
    item = request.get_json(silent=True) or {}
    item["id"] = item.get("id") or uuid.uuid4().hex[:12]
    item["ts"] = item.get("ts") or int(time.time() * 1000)
    records.insert(0, item)
    records = records[:120]
    with open(path, "w", encoding="utf-8") as f:
        json.dump(records, f, ensure_ascii=False, indent=2)
    return _json({"ok": True, "id": item["id"]})


if __name__ == "__main__":
    ensure_demo_docx()
    print(u"论小研已启动：http://127.0.0.1:8080")
    app.run(host="127.0.0.1", port=8080, debug=False, threaded=True)
